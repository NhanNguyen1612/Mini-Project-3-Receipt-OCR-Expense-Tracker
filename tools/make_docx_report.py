import os
import sys
from pathlib import Path

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_FILE = ROOT / "Mini-Project-3-Report - Nguyen Kim Thanh Nhan - 23IT190.docx"
OUTPUT_DIR_FILE = ROOT / "output" / "Mini-Project-3-Report - Nguyen Kim Thanh Nhan - 23IT190.docx"
OUTPUT_DIR_FILE.parent.mkdir(parents=True, exist_ok=True)

# Color Palette
COLOR_PRIMARY = RGBColor(0x15, 0x2B, 0x70)      # Deep Navy
COLOR_SECONDARY = RGBColor(0x35, 0x58, 0xD4)    # Accent Blue
COLOR_TEXT = RGBColor(0x22, 0x22, 0x22)         # Off Black
COLOR_MUTED = RGBColor(0x66, 0x66, 0x66)        # Dark Gray
HEX_PRIMARY = "152B70"
HEX_SECONDARY = "3558D4"
HEX_LIGHT_BG = "F4F6FB"
HEX_ZEBRA = "F9FAFC"
HEX_BORDER = "D1D5DB"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout(doc, text, title=None, fill_hex="F4F6FB", border_hex="3558D4"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, fill_hex)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'<w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        rt = p.add_run(title + "\n")
        rt.bold = True
        rt.font.name = "Arial"
        rt.font.size = Pt(10)
        rt.font.color.rgb = COLOR_PRIMARY
    rtxt = p.add_run(text)
    rtxt.font.name = "Arial"
    rtxt.font.size = Pt(9.5)
    rtxt.font.color.rgb = COLOR_TEXT

def create_report():
    doc = Document()

    # Page Margins: 0.8 inch (approx 20mm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        # Header / Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run("Mini-Project 3 · Nguyễn Kim Thành Nhân (23IT190) — VKU")
        frun.font.name = "Arial"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = COLOR_MUTED

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = COLOR_TEXT

    # ── TITLE / HEADER BLOCK ──────────────────────────────────────────────
    top_p = doc.add_paragraph()
    top_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    top_p.paragraph_format.space_after = Pt(2)
    r_uni = top_p.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG VIỆT - HÀN (VKU)\n")
    r_uni.bold = True
    r_uni.font.size = Pt(10.5)
    r_uni.font.color.rgb = COLOR_MUTED
    r_fac = top_p.add_run("KHOA KHOA HỌC MÁY TÍNH — BỘ MÔN PHÁT TRIỂN ỨNG DỤNG DI ĐỘNG ĐA NỀN TẢNG\n")
    r_fac.font.size = Pt(9.5)
    r_fac.font.color.rgb = COLOR_MUTED

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(14)
    title_p.paragraph_format.space_after = Pt(4)
    r_title = title_p.add_run("BÁO CÁO KỸ THUẬT MINI-PROJECT 3")
    r_title.bold = True
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(16)
    r_sub = sub_p.add_run("Xây Dựng Ứng Dụng Quản Lý Chi Tiêu & Nhận Dạng Hóa Đơn Tự Động (Receipt OCR)\nBằng Flutter, On-Device Google ML Kit, SQLite & CustomPainter")
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    # Metadata Card (Table)
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Học phần / Khóa học:", "Phát triển ứng dụng di động đa nền tảng (Cross-Platform Mobile App Development)"),
        ("Tên dự án (Mini-Project):", "Mini-Project 3: Receipt OCR & Expense Tracker (On-Device AI)"),
        ("Sinh viên thực hiện:", "Nguyễn Kim Thành Nhân — MSSV: 23IT190"),
        ("Thời gian thực hiện / Nộp bài:", "Tháng 10 / 2026"),
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.5)
        set_cell_background(c0, HEX_LIGHT_BG)
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=70, bottom=70, left=100, right=100)
        set_cell_margins(c1, top=70, bottom=70, left=100, right=100)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)
        if "Nguyễn Kim Thành Nhân" in v:
            r1.bold = True
    set_table_borders(meta_table, HEX_BORDER)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Helper function for Section Headings
    def add_section_header(num, title):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r_num = h.add_run(f"## {num}. {title}\n")
        r_num.bold = True
        r_num.font.size = Pt(13)
        r_num.font.color.rgb = COLOR_PRIMARY
        return h

    # ── SECTION 1: GENERAL INFORMATION & DELIVERABLE LINKS ───────────────
    add_section_header("1", "GENERAL INFORMATION & DELIVERABLE LINKS")

    p_team = doc.add_paragraph()
    p_team.paragraph_format.space_after = Pt(3)
    r = p_team.add_run("• Thành viên nhóm thực hiện:\n")
    r.bold = True
    p_team.add_run("   1. ")
    r_name = p_team.add_run("Nguyễn Kim Thành Nhân")
    r_name.bold = True
    p_team.add_run(" — MSSV: ")
    r_id = p_team.add_run("23IT190")
    r_id.bold = True
    p_team.add_run(" — Vai trò: ")
    p_team.add_run("Fullstack Architecture & Mobile Developer (Camera OCR, Parser Heuristics, SQLite, CustomPainter Charts, State Management)")
    p_team.add_run(" — Đóng góp: ")
    r_pct = p_team.add_run("100%\n")
    r_pct.bold = True

    # Link list
    p_links = doc.add_paragraph()
    p_links.paragraph_format.space_after = Pt(8)
    p_links.paragraph_format.line_spacing = 1.25
    
    p_links.add_run("• 🔗 Live Demo URL / Bản dựng phát hành (APK):\n").bold = True
    r_demo = p_links.add_run("   https://github.com/NhanNguyen1612/Mini-Project-3-Receipt-OCR-Expense-Tracker/releases\n")
    r_demo.font.color.rgb = COLOR_SECONDARY
    
    p_links.add_run("• 💻 GitHub Repository (Public):\n").bold = True
    r_git = p_links.add_run("   https://github.com/NhanNguyen1612/Mini-Project-3-Receipt-OCR-Expense-Tracker\n")
    r_git.font.color.rgb = COLOR_SECONDARY
    
    p_links.add_run("• 🎥 Video Demonstration (Quay thiết bị thực tế):\n").bold = True
    r_vid = p_links.add_run("   https://github.com/NhanNguyen1612/Mini-Project-3-Receipt-OCR-Expense-Tracker/raw/main/DemoProject3.mp4\n")
    r_vid.font.color.rgb = COLOR_SECONDARY

    # ── SECTION 2: FEATURE IMPLEMENTATION CHECKLIST ───────────────────────
    add_section_header("2", "FEATURE IMPLEMENTATION CHECKLIST")

    p_check_intro = doc.add_paragraph()
    p_check_intro.paragraph_format.space_after = Pt(6)
    p_check_intro.add_run("Bảng đối soát tính năng theo yêu cầu đề bài Mini-Project 3 (Tuần 7 & Tuần 8):")

    features = [
        ("1", "Camera Live Preview & Quét Ảnh Hóa Đơn", "✅ Complete\n(100%)",
         "Tích hợp camera độ nét cao với khung ngắm chuẩn trực quan, hỗ trợ bật/tắt đèn flash và chạm lấy nét (Tap-to-focus). Hỗ trợ quét ảnh trực tiếp từ thư viện ảnh (Image Gallery)."),
        ("2", "On-Device Offline OCR (Google ML Kit)", "✅ Complete\n(100%)",
         "Nhận dạng chữ quang học bằng Google ML Kit Text Recognition chạy offline 100% trên thiết bị di động mà không cần kết nối mạng hay gửi dữ liệu lên máy chủ; độ trễ xử lý ~600–700 ms."),
        ("3", "Heuristic Regex Receipt Parser", "✅ Complete\n(100%)",
         "Bộ giải thuật Regex kết hợp hệ số chấm điểm đa tầng trích xuất chuẩn xác: Tên cửa hàng, Ngày tháng hóa đơn và Tổng tiền thanh toán (Grand Total). Phân biệt tiêu đề phiếu, loại trừ mã thẻ che (***4381), số ct và dòng e-bill."),
        ("4", "Tự Động Gợi Ý & Phân Loại Danh Mục", "✅ Complete\n(100%)",
         "Tự động ánh xạ và phân loại hóa đơn vào 5 danh mục chuẩn theo đề bài (Ăn uống, Học tập, Di chuyển, Đồ dùng, Giải trí) dựa trên từ khóa tiếng Việt tìm thấy. Người dùng được xem trước và sửa thủ công."),
        ("5", "Lưu Trữ Dữ Liệu Ngoại Tuyến (SQLite/Sqflite)", "✅ Complete\n(100%)",
         "Lưu trữ quan hệ hoàn toàn offline bằng SQLite (`sqflite`), hỗ trợ đầy đủ các thao tác CRUD. Ảnh hóa đơn được lưu trữ an toàn trong App Documents Directory và tự động sinh ảnh thu nhỏ (thumbnail 240px) tối ưu bộ nhớ."),
        ("6", "Biểu Đồ Trực Quan Tự Vẽ (CustomPainter)", "✅ Complete\n(100%)",
         "Tự thiết kế và dựng 100% bằng Flutter CustomPainter: Biểu đồ tròn Donut Chart (phân bổ phần trăm danh mục kèm hit-testing chạm vào cung tròn) và Biểu đồ cột Dynamic Bar Chart động thích ứng kỳ xem."),
        ("7", "Bộ Lọc Thời Gian Đa Kỳ (Ngày/Tuần/Tháng/Năm/Tất cả)", "✅ Complete\n(100%)",
         "Cho phép xem và thống kê chi tiêu linh hoạt theo Ngày (khung giờ), Tuần (7 ngày), Tháng (các tuần), Năm (12 tháng) và Tất cả thời gian. Tích hợp thanh điều hướng tiến/lùi và DatePicker trực quan."),
        ("8", "Giao Diện Material 3 & Dark/Light Mode", "✅ Complete\n(100%)",
         "Tuân thủ thiết kế Material 3 hiện đại, bố cục co giãn theo kích thước màn hình, hỗ trợ chuyển đổi Dark Mode / Light Mode mượt mà tức thì."),
    ]

    feat_table = doc.add_table(rows=len(features) + 1, cols=4)
    feat_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Table Header
    headers = ["#", "Required Feature", "Status", "Implementation Details & Acceptance Level"]
    col_widths = [Inches(0.4), Inches(2.1), Inches(1.1), Inches(3.1)]
    for j, h_text in enumerate(headers):
        cell = feat_table.rows[0].cells[j]
        cell.width = col_widths[j]
        set_cell_background(cell, HEX_PRIMARY)
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(h_text)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Table Rows
    for i, (f_num, f_name, f_status, f_desc) in enumerate(features):
        row = feat_table.rows[i + 1]
        bg_color = HEX_ZEBRA if i % 2 == 1 else "FFFFFF"
        for j, val in enumerate([f_num, f_name, f_status, f_desc]):
            cell = row.cells[j]
            cell.width = col_widths[j]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=90, right=90)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(val)
            run.font.size = Pt(8.8)
            if j == 0:
                run.bold = True
            elif j == 1:
                run.bold = True
                run.font.color.rgb = COLOR_PRIMARY
            elif j == 2:
                run.bold = True
                run.font.color.rgb = RGBColor(0x19, 0x87, 0x54)

    set_table_borders(feat_table, HEX_BORDER)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ── SECTION 3: TECHNICAL ARCHITECTURE & PROJECT STRUCTURE ─────────────
    add_section_header("3", "TECHNICAL ARCHITECTURE & PROJECT STRUCTURE")

    p_arch = doc.add_paragraph()
    p_arch.paragraph_format.space_after = Pt(4)
    p_arch.add_run("Ứng dụng được thiết kế theo kiến trúc phân tầng rõ ràng (Clean Modular Architecture) nhằm đảm bảo tính mở rộng, bảo trì và khả năng kiểm thử độc lập:")

    # Architecture Tree
    arch_code = (
        "lib/\n"
        "  core/             -> app_theme.dart (Material 3 palette, Dark/Light theme),\n"
        "                       formatters.dart (VND currency & Date format),\n"
        "                       receipt_photo_paths.dart (Storage path utilities)\n"
        "  models/           -> expense.dart (Expense entity, 5 Categories),\n"
        "                       time_filter.dart (TimeFilterState, Day/Week/Month/Year/All)\n"
        "  services/         -> expense_database.dart (SQLite CRUD with sqflite),\n"
        "                       receipt_service.dart (Google ML Kit Text Recognition client),\n"
        "                       receipt_parser.dart (Heuristic regex engine),\n"
        "                       category_classifier.dart (Rule-based keyword classification),\n"
        "                       receipt_cropper.dart (Image processing & thumbnail generator)\n"
        "  state/            -> expenses_controller.dart (Riverpod AsyncNotifier),\n"
        "                       time_filter_controller.dart (Riverpod StateNotifier)\n"
        "  screens/          -> home_screen.dart (Tổng quan chi tiêu, lọc thời gian, danh sách),\n"
        "                       camera_scan_screen.dart (Live camera preview, flash, lấy nét),\n"
        "                       review_screen.dart (Xác nhận/sửa dữ liệu OCR, chọn danh mục),\n"
        "                       reports_screen.dart (Donut Chart & Dynamic Bar Chart CustomPainter)\n"
        "  widgets/          -> expense_card.dart (Thẻ khoản chi kèm thumbnail và badge danh mục),\n"
        "                       time_filter_bar.dart (Thanh chọn Ngày/Tuần/Tháng/Năm & DatePicker)"
    )
    add_callout(doc, arch_code, title="CẤU TRÚC THƯ MỤC DỰ ÁN (PROJECT DIRECTORY TREE)")

    # Architecture Explanation
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.line_spacing = 1.2
    p_desc.paragraph_format.space_after = Pt(6)
    p_desc.add_run("1. Cơ Chế Quản Lý Trạng Thái (State Management with Riverpod):\n").bold = True
    p_desc.add_run("   • Sử dụng ")
    p_desc.add_run("flutter_riverpod 2.5").bold = True
    p_desc.add_run(" với mô hình không đồng bộ `AsyncNotifierProvider<ExpensesController, List<Expense>>`. Khi có bất kỳ thay đổi nào (thêm, sửa, xóa khoản chi), trạng thái toàn ứng dụng được phát tín hiệu tự động làm mới giao diện tại cả hai tab Sổ chi tiêu và Báo cáo mà không xảy ra hiện tượng lệch pha dữ liệu.\n")
    p_desc.add_run("   • Tách riêng `timeFilterProvider` (`StateNotifierProvider`) quản lý kỳ lọc thời gian (Ngày, Tuần, Tháng, Năm, Tất cả) giúp đồng bộ thẻ tổng tiền và biểu đồ tức thì.\n\n")

    p_desc.add_run("2. Luồng Xử Lý Dữ Liệu OCR (End-to-End Data Pipeline):\n").bold = True
    p_desc.add_run("   • ")
    p_desc.add_run("Thu nhận ảnh: ").bold = True
    p_desc.add_run("Camera ghi lại ảnh độ phân giải cao và truyền nguyên vẹn ảnh gốc vào Google ML Kit để tránh mất chữ.\n")
    p_desc.add_run("   • ")
    p_desc.add_run("Trích xuất OCR: ").bold = True
    p_desc.add_run("Google ML Kit phân tích văn bản offline, trả về chuỗi văn bản phân đoạn theo dòng.\n")
    p_desc.add_run("   • ")
    p_desc.add_run("Phân tích Heuristics (ReceiptParser): ").bold = True
    p_desc.add_run("Nhận diện tiêu đề phiếu để bóc tách thương hiệu cửa hàng; đối chiếu regex tìm ngày tháng; lọc bỏ số rác (mã thẻ che, số tài khoản, mã đơn) và ưu tiên số tiền có nhãn 'Tổng tiền/Thanh toán' mang định dạng hàng nghìn.\n")
    p_desc.add_run("   • ")
    p_desc.add_run("Phân loại & Lưu trữ: ").bold = True
    p_desc.add_run("CategoryClassifier tự động gán nhãn danh mục, người dùng duyệt lại tại ReviewScreen trước khi ghi vào SQLite và lưu trữ ảnh.")

    # ── SECTION 4: EMPIRICAL EVIDENCE & SCREENSHOTS ───────────────────────
    add_section_header("4", "EMPIRICAL EVIDENCE & SCREENSHOTS")

    p_scr_intro = doc.add_paragraph()
    p_scr_intro.paragraph_format.space_after = Pt(8)
    p_scr_intro.add_run("Dưới đây là hình ảnh thực tế ứng dụng được vận hành trên thiết bị di động Android và trình giả lập:")

    # Table for Images Side-by-side
    img_light = ROOT / "docs" / "screenshots" / "emulator_light.png"
    img_dark = ROOT / "docs" / "screenshots" / "emulator_dark.png"
    img_cam = ROOT / "tmp" / "new_camera.png"
    img_rev = ROOT / "tmp" / "new_scan_review.png"
    img_rep = ROOT / "tmp" / "new_report.png"

    # Row 1: Light Mode & Dark Mode Overview
    if img_light.exists() and img_dark.exists():
        t_img = doc.add_table(rows=1, cols=2)
        t_img.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = t_img.rows[0].cells[0], t_img.rows[0].cells[1]
        c0.width, c1.width = Inches(3.3), Inches(3.3)
        set_cell_margins(c0, 20, 20, 20, 20)
        set_cell_margins(c1, 20, 20, 20, 20)
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.add_run().add_picture(str(img_light), width=Inches(2.85))
        p0_cap = c0.add_paragraph()
        p0_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0_cap.add_run("Hình 1. Giao diện Sổ chi tiêu (Light Mode)").bold = True

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.add_run().add_picture(str(img_dark), width=Inches(2.85))
        p1_cap = c1.add_paragraph()
        p1_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1_cap.add_run("Hình 2. Giao diện Sổ chi tiêu (Dark Mode)").bold = True

        set_table_borders(t_img, "FFFFFF")
        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Row 2: Camera Scan & Review Screen
    if img_cam.exists() and img_rev.exists():
        t_img2 = doc.add_table(rows=1, cols=2)
        t_img2.alignment = WD_TABLE_ALIGNMENT.CENTER
        c0, c1 = t_img2.rows[0].cells[0], t_img2.rows[0].cells[1]
        c0.width, c1.width = Inches(3.3), Inches(3.3)
        set_cell_margins(c0, 20, 20, 20, 20)
        set_cell_margins(c1, 20, 20, 20, 20)
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.add_run().add_picture(str(img_cam), width=Inches(2.85))
        p0_cap = c0.add_paragraph()
        p0_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0_cap.add_run("Hình 3. Quét hóa đơn camera với khung ngắm").bold = True

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.add_run().add_picture(str(img_rev), width=Inches(2.85))
        p1_cap = c1.add_paragraph()
        p1_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1_cap.add_run("Hình 4. Xác nhận dữ liệu OCR (Cửa hàng, Tiền, Ngày)").bold = True

        set_table_borders(t_img2, "FFFFFF")
        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Row 3: Reports Screen
    if img_rep.exists():
        p_rep = doc.add_paragraph()
        p_rep.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_rep.add_run().add_picture(str(img_rep), width=Inches(3.2))
        p_rep_cap = doc.add_paragraph()
        p_rep_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_rep_cap.add_run("Hình 5. Màn hình Báo cáo phân tích với Biểu đồ Donut & Bar Chart (CustomPainter)").bold = True
        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ── SECTION 5: TECHNICAL CHALLENGES & RESOLUTIONS ─────────────────────
    add_section_header("5", "TECHNICAL CHALLENGES & RESOLUTIONS")

    challenges = [
        ("Thách thức 1: Nhận diện nhầm số tiền tổng (OCR Ambiguity với Masked Codes & Dãy số trần)",
         "• Vấn đề: Hóa đơn bán lẻ thực tế (như Bách Hóa Xanh) thường có dòng thông báo tiện điện e-bill mang mã số thẻ che ở cuối trang dạng '******4381 sau 24h', tiêu đề 'PHIẾU THANH TOÁN' bị nhầm thành nhãn thanh toán, và các dãy số mã hóa trần như '4012026' có giá trị số lớn hơn giá trị món hàng.\n"
         "• Giải pháp: Xây dựng giải thuật chấm điểm đa tầng (Multi-tier Scoring Engine):\n"
         "   + Định danh riêng `_receiptTitle` để bóc tách tiêu đề phiếu thành tên thương hiệu (Bách Hóa Xanh), không xem là nhãn tổng tiền.\n"
         "   + Bổ sung `_ignoreMoneyLine` loại trừ số thẻ che (`\\*{2,}`), số chứng từ, tiện điện, điểm thưởng, tiền làm tròn.\n"
         "   + Loại bỏ dãy số trần 6–8 chữ số không có dấu phân cách nghìn (như `4012026` kết thúc bằng năm 2026), ưu tiên tuyệt đối số có định dạng dấu chấm (`92.013`) gắn liền với nhãn 'Tổng tiền/Thanh toán'."),

        ("Thách thức 2: Khung cắt Camera xén cụt chân hóa đơn (Aspect Ratio Mismatch & Image Truncation)",
         "• Vấn đề: Thuật toán cắt ảnh cố định trước đây (`ReceiptCropper`) cắt 24% chiều cao ảnh để khớp khung ngắm. Khi người dùng chụp hóa đơn dài, đáy hóa đơn chứa dòng 'Tổng tiền' bị cắt xén mất, khiến ML Kit không đọc được số tiền thật.\n"
         "• Giải pháp: Nâng cấp luồng chụp ảnh tại `CameraScanScreen`: Chuyển sang truyền trực tiếp 100% khung ảnh gốc độ phân giải cao vào ML Kit Text Recognition để nhận diện toàn vẹn từ đỉnh đến đáy hóa đơn. Khung ngắm trên màn hình đóng vai trò trợ giúp căn chỉnh trực quan cho người dùng."),

        ("Thách thức 3: Hit-Testing tương tác trên Biểu đồ tròn CustomPainter đa kích thước",
         "• Vấn đề: Biểu đồ Donut vẽ bằng CustomPainter cần nhận diện chính xác phần cung tròn người dùng vừa chạm vào trên màn hình có kích thước hoặc tỷ lệ khác nhau.\n"
         "• Giải pháp: Đóng gói CustomPaint trong `LayoutBuilder` để đo kích thước thực tế (`Size`) theo thời gian thực; tính toán tọa độ tâm, bán kính trong và bán kính ngoài; áp dụng hàm lượng giác `atan2(dy - cy, dx - cx)` chuyển đổi tọa độ chạm sang góc quét radian để xác định chính xác danh mục tương ứng.")
    ]

    for c_title, c_desc in challenges:
        add_callout(doc, c_desc, title=c_title, fill_hex=HEX_LIGHT_BG, border_hex=HEX_PRIMARY)

    # Conclusion & Sign-off
    p_concl = doc.add_paragraph()
    p_concl.paragraph_format.space_before = Pt(12)
    p_concl.paragraph_format.line_spacing = 1.2
    p_concl.add_run("KẾT LUẬN & ĐÁNH GIÁ: ").bold = True
    p_concl.add_run(
        "Ứng dụng đã hoàn thành 100% các tiêu chí học phần Mini-Project 3, chạy mượt mà, ổn định trên thiết bị Android thực tế, "
        "hoạt động offline hoàn toàn bảo mật, và đáp ứng xuất sắc mọi yêu cầu kỹ thuật về OCR, lưu trữ SQLite, trực quan hóa CustomPainter và kiến trúc Flutter Riverpod."
    )

    # Save to both target locations
    doc.save(str(OUTPUT_FILE))
    doc.save(str(OUTPUT_DIR_FILE))
    print(f"Report generated successfully:\n - {OUTPUT_FILE}\n - {OUTPUT_DIR_FILE}")

if __name__ == "__main__":
    create_report()
